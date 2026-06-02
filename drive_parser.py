import queue
import re
import sys
import threading
import tkinter as tk
from dataclasses import astuple, dataclass
from tkinter import filedialog, messagebox, ttk

import customtkinter as ctk
import openpyxl
import openpyxl.utils
import requests
from openpyxl.styles import Alignment, Font, PatternFill

# ── Domain ────────────────────────────────────────────────────────────────────

COLUMNS = [
    ("folder",   "Folder",       200),
    ("name",     "Name",         260),
    ("mime",     "MIME Type",    200),
    ("size",     "Size (bytes)", 110),
    ("modified", "Modified",      90),
    ("id",       "File ID",      180),
]

FOLDER_MIME    = "application/vnd.google-apps.folder"
DRIVE_FILES_URL = "https://www.googleapis.com/drive/v3/files"
MOD = "Command" if sys.platform == "darwin" else "Control"


def _natural_key(s: str) -> list:
    """Split a string into text/int chunks so '2' < '11' (natural sort)."""
    return [int(c) if c.isdigit() else c.lower() for c in re.split(r"(\d+)", s)]


@dataclass
class DriveFile:
    folder:   str
    name:     str
    mime:     str
    size:     str
    modified: str
    id:       str


# ── API (runs in background thread) ──────────────────────────────────────────

def scan_folder(token: str, folder_id: str, q: queue.Queue, stop: threading.Event) -> None:
    try:
        _recurse(token, folder_id, q, stop)
        if not stop.is_set():
            q.put(("done", None))
    except Exception as exc:
        q.put(("error", str(exc)))


def _recurse(token: str, folder_id: str, q: queue.Queue, stop: threading.Event, path: str = "/") -> None:
    page_token = None
    headers    = {"Authorization": f"Bearer {token}"}

    while not stop.is_set():
        params = {
            "q":                       f"'{folder_id}' in parents and trashed=false",
            "fields":                  "nextPageToken, files(id, name, mimeType, size, modifiedTime)",
            "pageSize":                100,
            "includeItemsFromAllDrives": "true",
            "supportsAllDrives":       "true",
        }
        if page_token:
            params["pageToken"] = page_token

        resp = requests.get(DRIVE_FILES_URL, headers=headers, params=params, timeout=30)

        if resp.status_code != 200:
            msg = resp.json().get("error", {}).get("message", resp.text)
            raise RuntimeError(f"API error {resp.status_code}: {msg}")

        data = resp.json()

        for item in data.get("files", []):
            if stop.is_set():
                return
            if item.get("mimeType") == FOLDER_MIME:
                child_path = f"{path}{item['name']}/"
                q.put(("status", f"Scanning: {child_path}"))
                _recurse(token, item["id"], q, stop, child_path)
            else:
                q.put(("file", DriveFile(
                    folder   = path,
                    name     = item.get("name", ""),
                    mime     = item.get("mimeType", ""),
                    size     = item.get("size", "—"),
                    modified = item.get("modifiedTime", "")[:10],
                    id       = item["id"],
                )))

        page_token = data.get("nextPageToken")
        if not page_token:
            break


# ── GUI ───────────────────────────────────────────────────────────────────────

ctk.set_appearance_mode("dark")
ctk.set_default_color_theme("blue")

# Treeview colours (must match CTk dark palette)
TC = {
    "bg":      "#1e1f22",
    "even":    "#2b2d30",
    "odd":     "#252528",
    "head_bg": "#18191c",
    "head_fg": "#7a8499",
    "sel_bg":  "#1f538d",
    "sel_fg":  "#ffffff",
    "fg":      "#d4d4d4",
}


class App(ctk.CTk):
    def __init__(self) -> None:
        super().__init__()
        self.title("Google Drive Parser")
        self.geometry("1060x700")
        self.minsize(800, 520)

        self._files:    list[DriveFile] = []
        self._visible:  list[int]       = []   # indices into _files
        self._q:        queue.Queue     = queue.Queue()
        self._stop:     threading.Event = threading.Event()

        self._filter_var = tk.StringVar()
        self._filter_var.trace_add("write", lambda *_: self._apply_filter())
        self._regex_mode   = False
        self._strip_prefix = False

        self._apply_tree_style()
        self._build_toolbar()
        self._build_search()
        self._build_table()
        self._build_statusbar()
        self._bind_clipboard()

    # ── Treeview theme ────────────────────────────────────────────────────────

    def _apply_tree_style(self) -> None:
        s = ttk.Style(self)
        s.theme_use("clam")
        s.configure("Treeview",
            background=TC["even"], foreground=TC["fg"],
            fieldbackground=TC["even"], rowheight=28,
            borderwidth=0, relief="flat",
            font=("Helvetica Neue", 11),
        )
        s.configure("Treeview.Heading",
            background=TC["head_bg"], foreground=TC["head_fg"],
            font=("Helvetica Neue", 11, "bold"),
            borderwidth=0, relief="flat", padding=(8, 8),
        )
        s.map("Treeview",
            background=[("selected", TC["sel_bg"])],
            foreground=[("selected", TC["sel_fg"])],
        )
        s.map("Treeview.Heading",
            background=[("active", "#202124")],
            relief=[("active", "flat")],
        )
        s.configure("Vertical.TScrollbar",
            troughcolor=TC["bg"], background="#3c3f41", borderwidth=0, arrowsize=14)
        s.configure("Horizontal.TScrollbar",
            troughcolor=TC["bg"], background="#3c3f41", borderwidth=0, arrowsize=14)

    # ── Layout ────────────────────────────────────────────────────────────────

    def _build_toolbar(self) -> None:
        bar = ctk.CTkFrame(self, corner_radius=0, fg_color="#18191c")
        bar.pack(fill="x", padx=0, pady=0)

        inner = ctk.CTkFrame(bar, corner_radius=0, fg_color="transparent")
        inner.pack(fill="x", padx=16, pady=12)
        inner.columnconfigure(1, weight=1)

        ctk.CTkLabel(inner, text="Access Token", font=("Helvetica Neue", 12),
                     text_color="#7a8499").grid(row=0, column=0, sticky="w")
        self._token_entry = ctk.CTkEntry(
            inner, placeholder_text="ya29.xxx…", show="•",
            font=("Helvetica Neue", 12), height=36,
            fg_color="#2b2d30", border_color="#43454a", border_width=1,
            text_color="#d4d4d4",
        )
        self._token_entry.grid(row=0, column=1, padx=(10, 20), sticky="ew")

        ctk.CTkLabel(inner, text="Folder ID", font=("Helvetica Neue", 12),
                     text_color="#7a8499").grid(row=0, column=2, sticky="w")
        self._folder_entry = ctk.CTkEntry(
            inner, placeholder_text="1BxiM…", width=220,
            font=("Helvetica Neue", 12), height=36,
            fg_color="#2b2d30", border_color="#43454a", border_width=1,
            text_color="#d4d4d4",
        )
        self._folder_entry.grid(row=0, column=3, padx=(10, 20))

        self._btn_start = ctk.CTkButton(
            inner, text="▶  Start", command=self._on_start,
            width=110, height=36, font=("Helvetica Neue", 12, "bold"),
            corner_radius=8, fg_color="#1f538d", hover_color="#2563a8",
        )
        self._btn_start.grid(row=0, column=4, padx=(0, 8))

        self._btn_stop = ctk.CTkButton(
            inner, text="■  Stop", command=self._on_stop, state="disabled",
            width=90, height=36, font=("Helvetica Neue", 12, "bold"),
            corner_radius=8, fg_color="#962d22", hover_color="#b33628",
        )
        self._btn_stop.grid(row=0, column=5)

    def _build_search(self) -> None:
        bar = ctk.CTkFrame(self, corner_radius=0, fg_color="#1e1f22")
        bar.pack(fill="x", padx=16, pady=(8, 4))

        self._filter_entry = ctk.CTkEntry(
            bar, textvariable=self._filter_var,
            placeholder_text="🔍  Filter by name…",
            font=("Helvetica Neue", 12), height=34,
            fg_color="#2b2d30", border_color="#3c3f41", border_width=1,
            text_color="#d4d4d4",
        )
        self._filter_entry.pack(side="left", fill="x", expand=True, padx=(0, 8))

        self._btn_regex = ctk.CTkButton(
            bar, text=".*", width=42, height=34,
            font=("Helvetica Neue", 13, "bold"),
            corner_radius=6,
            fg_color="#2b2d30", hover_color="#3a3d42",
            border_width=1, border_color="#3c3f41",
            text_color="#7a8499",
            command=self._toggle_regex,
        )
        self._btn_regex.pack(side="left", padx=(0, 4))

        self._btn_strip = ctk.CTkButton(
            bar, text="1.", width=42, height=34,
            font=("Helvetica Neue", 13, "bold"),
            corner_radius=6,
            fg_color="#2b2d30", hover_color="#3a3d42",
            border_width=1, border_color="#3c3f41",
            text_color="#7a8499",
            command=self._toggle_strip_prefix,
        )
        self._btn_strip.pack(side="left")

    def _build_table(self) -> None:
        frame = ctk.CTkFrame(self, corner_radius=8, fg_color=TC["bg"])
        frame.pack(fill="both", expand=True, padx=16, pady=(0, 4))

        col_ids = [c[0] for c in COLUMNS]
        self._tree = ttk.Treeview(frame, columns=col_ids, show="headings", selectmode="extended")
        for key, label, width in COLUMNS:
            self._tree.heading(key, text=label, command=lambda k=key: self._sort_by(k))
            self._tree.column(key, width=width, anchor="w", minwidth=60)

        vsb = ttk.Scrollbar(frame, orient="vertical",   command=self._tree.yview)
        hsb = ttk.Scrollbar(frame, orient="horizontal", command=self._tree.xview)
        self._tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        vsb.pack(side="right",  fill="y")
        hsb.pack(side="bottom", fill="x")
        self._tree.pack(fill="both", expand=True)

        self._tree.tag_configure("even", background=TC["even"])
        self._tree.tag_configure("odd",  background=TC["odd"])

        self._copy_cols: set[str] = set()
        for key, label, _ in COLUMNS:
            self._tree.heading(key, text=label,
                               command=lambda k=key: self._sort_by(k))
        self._tree.bind("<Shift-Button-1>",          self._toggle_copy_col)
        self._tree.bind(f"<{MOD}-Button-1>",         self._cmd_click)
        self._tree.bind(f"<{MOD}-c>",                self._copy_selection)

    def _build_statusbar(self) -> None:
        bar = ctk.CTkFrame(self, corner_radius=0, fg_color="#18191c", height=48)
        bar.pack(fill="x", side="bottom")
        bar.pack_propagate(False)

        left = ctk.CTkFrame(bar, corner_radius=0, fg_color="transparent")
        left.pack(side="left", fill="y", padx=16)

        self._status_label = ctk.CTkLabel(
            left, text="Enter token and folder ID, then press Start.",
            font=("Helvetica Neue", 11), text_color="#6b7280", anchor="w",
        )
        self._status_label.pack(side="left", fill="y")

        self._count_label = ctk.CTkLabel(
            left, text="", font=("Helvetica Neue", 11, "bold"),
            text_color="#4a9eda", anchor="w",
        )
        self._count_label.pack(side="left", padx=(12, 0), fill="y")

        right = ctk.CTkFrame(bar, corner_radius=0, fg_color="transparent")
        right.pack(side="right", fill="y", padx=16)

        self._progress = ctk.CTkProgressBar(
            right, mode="indeterminate", width=120, height=6,
            progress_color="#1f538d", fg_color="#3c3f41",
        )

        ctk.CTkButton(
            right, text="Export Excel", command=self._export_xlsx,
            width=120, height=32, font=("Helvetica Neue", 11, "bold"),
            corner_radius=6, fg_color="#1a6b35", hover_color="#218a43",
        ).pack(side="right", pady=8)

    # ── Actions ───────────────────────────────────────────────────────────────

    def _on_start(self) -> None:
        token     = self._token_entry.get().strip()
        folder_id = self._folder_entry.get().strip()

        if not token or not folder_id:
            messagebox.showwarning("Missing input", "Fill in Access Token and Folder ID.")
            return

        self._files.clear()
        self._visible.clear()
        self._tree.delete(*self._tree.get_children())
        self._stop.clear()
        self._set_busy(True)
        self._set_status("Starting…")

        threading.Thread(
            target=scan_folder,
            args=(token, folder_id, self._q, self._stop),
            daemon=True,
        ).start()
        self._poll()

    def _on_stop(self) -> None:
        self._stop.set()

    def _export_xlsx(self) -> None:
        if not self._files:
            messagebox.showinfo("No data", "Run a search first.")
            return

        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel file", "*.xlsx")],
            initialfile="drive_files.xlsx",
        )
        if not path:
            return

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Drive Files"

        header_fill = PatternFill("solid", fgColor="1F538D")
        header_font = Font(bold=True, color="FFFFFF", size=11, name="Calibri")
        center      = Alignment(vertical="center", horizontal="left")

        ws.append([label for _, label, _ in COLUMNS])
        ws.row_dimensions[1].height = 28
        for cell in ws[1]:
            cell.fill = header_fill; cell.font = header_font; cell.alignment = center

        ws.freeze_panes = "A2"

        fill_even = PatternFill("solid", fgColor="FFFFFF")
        fill_odd  = PatternFill("solid", fgColor="F0F4F8")
        body_font = Font(size=10, name="Calibri")

        for i, f in enumerate(self._files):
            vals = list(astuple(f))
            try:
                vals[3] = int(vals[3])
            except (ValueError, TypeError):
                pass
            ws.append(vals)
            ws.row_dimensions[i + 2].height = 20
            row_fill = fill_even if i % 2 == 0 else fill_odd
            for cell in ws[i + 2]:
                cell.fill = row_fill; cell.font = body_font; cell.alignment = center

        for col_idx, (_, label, _) in enumerate(COLUMNS, start=1):
            letter  = openpyxl.utils.get_column_letter(col_idx)
            max_len = len(label)
            for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                for cell in row:
                    if cell.value is not None:
                        max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[letter].width = min(max_len + 4, 60)

        wb.save(path)
        messagebox.showinfo("Saved", f"Exported to:\n{path}")

    # ── Queue polling (main thread only) ─────────────────────────────────────

    def _poll(self) -> None:
        try:
            while True:
                kind, payload = self._q.get_nowait()
                if kind == "file":
                    self._ingest(payload)
                elif kind == "status":
                    self._set_status(payload)
                elif kind == "done":
                    self._set_status("Done")
                    self._update_count()
                    self._set_busy(False)
                    return
                elif kind == "error":
                    messagebox.showerror("Error", payload)
                    self._set_status("Error")
                    self._set_busy(False)
                    return
        except queue.Empty:
            pass

        if self._stop.is_set():
            self._set_status("Stopped")
            self._update_count()
            self._set_busy(False)
            return

        self.after(100, self._poll)

    # ── Table helpers ─────────────────────────────────────────────────────────

    def _ingest(self, f: DriveFile) -> None:
        self._files.append(f)
        raw = self._filter_var.get()
        pattern = None
        if raw and self._regex_mode:
            try:
                pattern = re.compile(raw, re.IGNORECASE)
            except re.error:
                pattern = None
        if not raw or self._matches(f, raw.lower(), pattern):
            self._visible.append(len(self._files) - 1)
            self._insert_row(f, len(self._visible) - 1)
        if len(self._files) % 50 == 0:
            self._update_count()

    def _insert_row(self, f: DriveFile, display_idx: int) -> None:
        tag = "even" if display_idx % 2 == 0 else "odd"
        self._tree.insert("", "end", values=astuple(f), tags=(tag,))

    def _toggle_regex(self) -> None:
        self._regex_mode = not self._regex_mode
        if self._regex_mode:
            self._btn_regex.configure(
                fg_color="#1f538d", hover_color="#2563a8",
                border_color="#1f538d", text_color="#ffffff",
            )
        else:
            self._btn_regex.configure(
                fg_color="#2b2d30", hover_color="#3a3d42",
                border_color="#3c3f41", text_color="#7a8499",
            )
        self._apply_filter()

    def _toggle_strip_prefix(self) -> None:
        self._strip_prefix = not self._strip_prefix
        if self._strip_prefix:
            self._btn_strip.configure(
                fg_color="#1f538d", hover_color="#2563a8",
                border_color="#1f538d", text_color="#ffffff",
            )
        else:
            self._btn_strip.configure(
                fg_color="#2b2d30", hover_color="#3a3d42",
                border_color="#3c3f41", text_color="#7a8499",
            )

    def _apply_filter(self) -> None:
        raw = self._filter_var.get()
        pattern = None

        if raw:
            if self._regex_mode:
                try:
                    pattern = re.compile(raw, re.IGNORECASE)
                    self._filter_entry.configure(border_color="#3c3f41")
                except re.error:
                    self._filter_entry.configure(border_color="#962d22")
                    return  # invalid regex — keep current view
            else:
                self._filter_entry.configure(border_color="#3c3f41")

        self._tree.delete(*self._tree.get_children())
        self._visible = [
            i for i, f in enumerate(self._files)
            if not raw or self._matches(f, raw.lower(), pattern)
        ]
        for display_idx, file_idx in enumerate(self._visible):
            self._insert_row(self._files[file_idx], display_idx)
        self._update_count()

    @staticmethod
    def _matches(f: DriveFile, q: str, pattern: re.Pattern | None) -> bool:
        if pattern:
            return bool(pattern.search(f.name))
        return q in f.name.lower()

    def _cmd_click(self, event: tk.Event) -> str:
        row = self._tree.identify_row(event.y)
        if row:
            sel = list(self._tree.selection())
            if row in sel:
                sel.remove(row)
            else:
                sel.append(row)
            self._tree.selection_set(sel)
        return "break"

    def _toggle_copy_col(self, event: tk.Event) -> str | None:
        if self._tree.identify_region(event.x, event.y) != "heading":
            return None  # let Shift+click row selection pass through
        col = self._tree.identify_column(event.x)
        if not col:
            return None
        key = COLUMNS[int(col[1:]) - 1][0]
        label = COLUMNS[int(col[1:]) - 1][1]
        if key in self._copy_cols:
            self._copy_cols.discard(key)
            self._tree.heading(key, text=label)
        else:
            self._copy_cols.add(key)
            self._tree.heading(key, text=f"· {label}")
        return "break"

    def _copy_selection(self, _=None) -> str:
        selected = self._tree.selection()
        if not selected or not self._copy_cols:
            return "break"
        col_indices = [i for i, (k, _, _) in enumerate(COLUMNS) if k in self._copy_cols]
        rows = []
        for iid in selected:
            vals = self._tree.item(iid, "values")
            parts = []
            for i in col_indices:
                val = str(vals[i])
                if self._strip_prefix:
                    val = re.sub(r'^(?:v\d+(?:\.\d+)*\.?|\d+(?:\.\d+)*[\.\)\-:])\s+', '', val)
                    val = re.sub(r'\.[a-zA-Z][a-zA-Z0-9]{0,4}$', '', val).strip()
                parts.append(val)
            rows.append("\t".join(parts))
        self.clipboard_clear()
        self.clipboard_append("\n".join(rows))
        return "break"

    def _repopulate(self) -> None:
        self._apply_filter()

    def _sort_by(self, key: str) -> None:
        idx = [c[0] for c in COLUMNS].index(key)
        folder_idx = [c[0] for c in COLUMNS].index("folder")
        self._files.sort(key=lambda f: (
            _natural_key(astuple(f)[folder_idx]),
            _natural_key(str(astuple(f)[idx])),
        ))
        self._apply_filter()

    # ── Status helpers ────────────────────────────────────────────────────────

    def _set_status(self, msg: str) -> None:
        self._status_label.configure(text=msg)

    def _update_count(self) -> None:
        total   = len(self._files)
        visible = len(self._visible)
        if total == 0:
            self._count_label.configure(text="")
        elif visible == total:
            self._count_label.configure(text=f"{total} files")
        else:
            self._count_label.configure(text=f"{visible} / {total} files")

    def _set_busy(self, busy: bool) -> None:
        self._btn_start.configure(state="disabled" if busy else "normal")
        self._btn_stop.configure(state="normal"    if busy else "disabled")
        if busy:
            self._progress.pack(side="right", padx=(0, 12), pady=20)
            self._progress.start()
        else:
            self._progress.stop()
            self._progress.pack_forget()

    # ── Clipboard bindings (cross-platform) ──────────────────────────────────

    def _bind_clipboard(self) -> None:
        def paste(e):
            try:
                text = self.clipboard_get()
                try:
                    e.widget.delete("sel.first", "sel.last")
                except tk.TclError:
                    pass
                e.widget.insert("insert", text)
            except tk.TclError:
                pass
            return "break"

        def copy(e):
            try:
                self.clipboard_clear()
                self.clipboard_append(e.widget.selection_get())
            except tk.TclError:
                pass
            return "break"

        def cut(e):
            try:
                text = e.widget.selection_get()
                self.clipboard_clear()
                self.clipboard_append(text)
                e.widget.delete("sel.first", "sel.last")
            except tk.TclError:
                pass
            return "break"

        def select_all(e):
            e.widget.select_range(0, "end")
            e.widget.icursor("end")
            return "break"

        for seq, fn in [(f"<{MOD}-v>", paste), (f"<{MOD}-c>", copy),
                        (f"<{MOD}-x>", cut),  (f"<{MOD}-a>", select_all)]:
            self.bind_class("Entry", seq, fn)


if __name__ == "__main__":
    App().mainloop()
